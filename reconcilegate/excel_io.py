"""Writes and reads the human-maintained-style Excel workbooks: a 2-row
merged header (ungrouped ID columns vertically merged across both header
rows, one grouped pair of columns horizontally merged under a category
label), one sheet per country, exactly the shape a plant analyst's manual
Excel log actually has. Reading it back is the real engineering problem:
openpyxl only ever populates the top-left cell of a merged range, every
other cell in that range is `None`, so the header has to be reconstructed
from `worksheet.merged_cells.ranges`, not read as two plain rows.
"""
from __future__ import annotations

from dataclasses import dataclass

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class SheetLayout:
    """Describes one workbook's header shape.

    id_columns: plain columns spanning both header rows (vertical merge),
                e.g. ["Order ID", "Plant", "Month"].
    group_label: the horizontal-merge category label over group_fields,
                 e.g. "Shipment". None if this workbook has no grouped pair
                 (complaint_tracker has none).
    group_fields: the sub-columns under group_label, e.g. ["Quantity", "Unit"].
    """

    id_columns: list
    group_label: str | None
    group_fields: list


def write_workbook(
    path: str, country_sheets: dict, layout: SheetLayout, header_label_override: dict | None = None
) -> None:
    """country_sheets: {country_code: DataFrame} with flat columns
    id_columns + group_fields (unprefixed).

    header_label_override: {country_code: [relabeled group_fields]}, used
    only to construct the deliberately-broken ``merged_header_misaligned``
    fixture: the row-2 labels under the merged group are written in this
    order while the data columns underneath stay in the layout's normal
    order, so the reconstructed header names the wrong data.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    n_id = len(layout.id_columns)
    n_group = len(layout.group_fields)
    header_label_override = header_label_override or {}

    for country, df in country_sheets.items():
        ws = wb.create_sheet(title=country)

        for i, name in enumerate(layout.id_columns):
            col = i + 1
            ws.cell(row=1, column=col, value=name)
            ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)

        if layout.group_label is not None:
            start_col = n_id + 1
            end_col = n_id + n_group
            ws.cell(row=1, column=start_col, value=layout.group_label)
            ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
            row2_labels = header_label_override.get(country, layout.group_fields)
            for j, field in enumerate(row2_labels):
                ws.cell(row=2, column=start_col + j, value=field)

        flat_columns = layout.id_columns + layout.group_fields
        for r, row in enumerate(df.itertuples(index=False), start=3):
            for c, value in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=value)

        for c in range(1, n_id + n_group + 1):
            ws.column_dimensions[get_column_letter(c)].width = 16

    wb.save(path)


def _reconstruct_header(ws) -> list:
    """Forward-fills merged category cells across their horizontal span,
    and takes the row-1 value for any column whose header is a vertical
    merge (row1 == row2's merge partner)."""
    max_col = ws.max_column
    row1 = [ws.cell(row=1, column=c).value for c in range(1, max_col + 1)]
    row2 = [ws.cell(row=2, column=c).value for c in range(1, max_col + 1)]

    # Forward-fill row1 across horizontal merges (None cells inside a merge).
    filled_row1 = []
    last = None
    for v in row1:
        if v is not None:
            last = v
        filled_row1.append(last)

    vertical_merge_cols = set()
    horizontal_merge_cols = set()
    for merged_range in ws.merged_cells.ranges:
        spans_rows = merged_range.max_row - merged_range.min_row
        spans_cols = merged_range.max_col - merged_range.min_col
        if merged_range.min_row == 1 and spans_rows >= 1 and spans_cols == 0:
            vertical_merge_cols.add(merged_range.min_col)
        if merged_range.min_row == 1 and spans_cols >= 1:
            for c in range(merged_range.min_col, merged_range.max_col + 1):
                horizontal_merge_cols.add(c)

    header = []
    for idx in range(max_col):
        col = idx + 1
        if col in vertical_merge_cols:
            header.append(filled_row1[idx])
        elif col in horizontal_merge_cols:
            header.append(f"{filled_row1[idx]}: {row2[idx]}")
        else:
            # Neither merge recorded (defensive fallback): prefer row2, else row1.
            header.append(row2[idx] if row2[idx] is not None else filled_row1[idx])
    return header


def read_workbook(path: str) -> dict:
    """Returns {sheet_name: DataFrame} with reconstructed column names,
    data rows starting at row 3."""
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header = _reconstruct_header(ws)
        rows = []
        for r in range(3, ws.max_row + 1):
            row = [ws.cell(row=r, column=c + 1).value for c in range(len(header))]
            if all(v is None for v in row):
                continue
            rows.append(row)
        out[sheet_name] = pd.DataFrame(rows, columns=header)
    return out
